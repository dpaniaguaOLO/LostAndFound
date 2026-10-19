"""Genera la presentación PowerPoint de Lost and Found para el supervisor.
Diseño con colores corporativos de OLO y capturas reales de la app.
Salida: docs/Lost_and_Found_OLO.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(__file__)
SHOTS = os.path.join(HERE, "..", "docs", "screenshots")
OUT = os.path.join(HERE, "..", "docs", "Lost_and_Found_OLO.pptx")

# --- Colores de marca OLO ---
OLO_GREEN = RGBColor(0x00, 0xA7, 0x85)   # verde de marca
OLO_DARK = RGBColor(0x00, 0x74, 0x5D)    # verde oscuro
INK = RGBColor(0x25, 0x1E, 0x1F)         # tinta (casi negro)
SLATE = RGBColor(0x64, 0x74, 0x8B)       # gris texto secundario
LIGHT = RGBColor(0xE6, 0xF7, 0xF2)       # verde muy claro (fondo acentos)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Lienzo 16:9
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Open Sans"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def fit_image(slide, path, box_x, box_y, box_w, box_h):
    """Inserta imagen ajustada dentro de un recuadro, centrada, sin deformar."""
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    box_ar = box_w / box_h
    if ar > box_ar:
        w = box_w
        h = int(box_w / ar)
    else:
        h = box_h
        w = int(box_h * ar)
    x = box_x + (box_w - w) // 2
    y = box_y + (box_h - h) // 2
    slide.shapes.add_picture(path, x, y, width=Emu(int(w)), height=Emu(int(h)))


# =========================================================================
# PORTADA
# =========================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
# Banda lateral verde
add_rect(s, 0, 0, Inches(0.35), SH, OLO_GREEN)
# Marca OLO
add_text(s, Inches(1), Inches(1.6), Inches(11), Inches(0.6),
         "OLO · Overseas Logistics Operations", 18, OLO_DARK, bold=True)
add_text(s, Inches(1), Inches(2.3), Inches(11.5), Inches(1.6),
         "Lost & Found", 54, INK, bold=True)
add_text(s, Inches(1), Inches(3.7), Inches(11), Inches(0.8),
         "Aplicación interna para gestionar objetos perdidos y encontrados",
         22, SLATE)
# Pastilla verde
add_rect(s, Inches(1), Inches(4.9), Inches(4.6), Inches(0.08), OLO_GREEN)
add_text(s, Inches(1), Inches(5.2), Inches(11), Inches(1.2),
         "Presentación funcional · Versión de desarrollo 0.3",
         14, SLATE)


# =========================================================================
# ¿QUÉ ES? (resumen en 3 puntos, sin captura)
# =========================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
add_rect(s, 0, 0, SW, Inches(1.2), OLO_GREEN)
add_text(s, Inches(0.7), Inches(0.28), Inches(12), Inches(0.7),
         "¿Qué es Lost & Found?", 30, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

puntos = [
    ("Para los trabajadores de OLO",
     "Una herramienta interna y sencilla para reportar y recuperar objetos dentro de las instalaciones."),
    ("Un muro común de objetos",
     "Lo que alguien encuentra se publica con foto, descripción, ubicación y fecha, visible para todo el personal."),
    ("Devoluciones ordenadas",
     "Quien reconoce un objeto lo reclama; el proceso queda registrado hasta la entrega."),
]
y = Inches(1.7)
for titulo, desc in puntos:
    add_rect(s, Inches(0.7), y + Inches(0.08), Inches(0.14), Inches(1.0), OLO_GREEN)
    add_text(s, Inches(1.05), y, Inches(11.4), Inches(0.5), titulo, 20, INK, bold=True)
    add_text(s, Inches(1.05), y + Inches(0.52), Inches(11.4), Inches(0.7), desc, 15, SLATE)
    y += Inches(1.6)


# =========================================================================
# Slides de funcionalidad: título + bullets + captura
# =========================================================================
features = [
    ("Acceso seguro", "01_login",
     ["Cada trabajador entra con su cuenta de OLO.",
      "Solo el personal autorizado ve el muro de objetos."]),
    ("Muro de objetos encontrados", "02_muro",
     ["Vista central con todo lo que se ha encontrado.",
      "Búsqueda y filtros por categoría y ubicación.",
      "Cada objeto muestra su estado (publicado, reclamado, devuelto)."]),
    ("Detalle del objeto", "03_detalle",
     ["Toda la información del objeto en una sola vista.",
      "Desde aquí se puede reclamar o marcar como devuelto."]),
    ("Publicar un hallazgo", "04_publicar",
     ["Formulario corto: foto, descripción, ubicación y fecha.",
      "Publicar un objeto toma menos de un minuto."]),
    ("Reportar objetos perdidos", "05_perdidos",
     ["El personal también reporta lo que ha perdido.",
      "Facilita relacionar perdidos con encontrados."]),
    ("Reclamaciones y devolución", "06_reclamaciones",
     ["Las solicitudes para recuperar un objeto se gestionan aquí.",
      "Se aprueban o rechazan y el objeto se marca como devuelto."]),
    ("Notificaciones", "07_notificaciones",
     ["Avisos dentro de la app sobre reclamaciones y devoluciones.",
      "El personal se entera sin salir de la herramienta."]),
    ("Panel de administración", "08_admin",
     ["Resumen del sistema con indicadores clave.",
      "Gestión de usuarios y permisos para encargados."]),
]

for titulo, shot, bullets in features:
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, SW, SH, WHITE)
    # Cabecera
    add_rect(s, 0, 0, SW, Inches(1.2), OLO_GREEN)
    add_text(s, Inches(0.7), Inches(0.28), Inches(12), Inches(0.7),
             titulo, 28, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    # Columna izquierda: bullets
    by = Inches(1.9)
    for b in bullets:
        add_rect(s, Inches(0.7), by + Inches(0.12), Inches(0.12), Inches(0.12), OLO_GREEN)
        add_text(s, Inches(1.0), by, Inches(3.6), Inches(0.9), b, 14, INK)
        by += Inches(1.0)
    # Columna derecha: captura dentro de un marco claro
    frame_x, frame_y = Inches(4.9), Inches(1.7)
    frame_w, frame_h = Inches(7.9), Inches(5.2)
    add_rect(s, frame_x, frame_y, frame_w, frame_h, LIGHT)
    pad = Inches(0.15)
    fit_image(s, os.path.join(SHOTS, f"{shot}.png"),
              frame_x + pad, frame_y + pad,
              frame_w - 2 * pad, frame_h - 2 * pad)


# =========================================================================
# CIERRE
# =========================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, INK)
add_rect(s, 0, Inches(3.5), SW, Inches(0.1), OLO_GREEN)
add_text(s, Inches(1), Inches(2.3), Inches(11.3), Inches(1.2),
         "Lost & Found", 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(3.8), Inches(11.3), Inches(0.8),
         "Una herramienta hecha para OLO", 20, OLO_GREEN, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.6), Inches(11.3), Inches(1.0),
         "Base funcional lista · próximos pasos: pruebas, mejoras y puesta en marcha interna",
         14, RGBColor(0xC0, 0xCB, 0xD3), align=PP_ALIGN.CENTER)

prs.save(OUT)
print("Presentación generada:", os.path.abspath(OUT))
print("Diapositivas:", len(prs.slides._sldIdLst))
